from fastapi import APIRouter, Request, Depends, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete, or_, and_  # 💡 and_ 추가됨
from app.core.database import get_db
from app.models.schedule import WorkSchedule
from app.models.event import Event
from datetime import datetime, time
from openpyxl import load_workbook
import calendar
import pandas as pd
import io
import re

router = APIRouter(prefix="/schedule", tags=["Schedule"])
templates = Jinja2Templates(directory="templates")

@router.get("/list", response_class=HTMLResponse)
async def schedule_list_page(request: Request):
    return templates.TemplateResponse("schedule_list.html", {"request": request})

@router.post("/upload")
async def upload_schedule_excel(file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    try:
        match = re.search(r'(\d{4})년\s*(\d{1,2})월', file.filename)
        if not match:
            raise HTTPException(status_code=400, detail="파일명 형식을 확인해주세요. (예: 2026년 2월)")

        target_year, target_month = int(match.group(1)), int(match.group(2))
        last_day = calendar.monthrange(target_year, target_month)[1]

        content = await file.read()
        wb = load_workbook(io.BytesIO(content), data_only=True)
        ws = wb["근무표_작성용"]

        NAME_COL = 3
        START_COL = 4
        START_ROW = 6
        END_ROW = 14
        DATE_ROW = 4

        records = []

        for r in range(START_ROW, END_ROW + 1):
            worker_name = str(ws.cell(row=r, column=NAME_COL).value or "").strip()
            if worker_name in ['0', 'None', '', 'nan']: continue

            for d in range(1, last_day + 1):
                curr_col = START_COL + d - 1
                curr_date = datetime(target_year, target_month, d).date()

                val = str(ws.cell(row=r, column=curr_col).value or "").strip()
                val = "" if val in ['0', '0.0', 'None', 'nan'] else val

                header_cell = ws.cell(row=DATE_ROW, column=curr_col)
                bg_color = header_cell.fill.start_color.index

                is_holiday = False
                if bg_color and str(bg_color) not in ["00000000", "FFFFFFFF", "None"]:
                    is_holiday = True

                records.append(WorkSchedule(
                    work_date=curr_date,
                    worker_name=worker_name,
                    shift_type=val,
                    is_holiday=is_holiday
                ))

        start_d = datetime(target_year, target_month, 1).date()
        end_d = datetime(target_year, target_month, last_day).date()
        await db.execute(delete(WorkSchedule).where(WorkSchedule.work_date.between(start_d, end_d)))
        db.add_all(records)
        await db.commit()

        return {"message": f"{target_year}년 {target_month}월 근무표가 좌표(C6:AH14) 기준으로 완벽히 반영되었습니다."}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"분석 실패: {str(e)}")

@router.get("/api/data")
async def get_schedule_data(year: int, month: int, db: AsyncSession = Depends(get_db)):
    start_date = datetime(year, month, 1).date()
    import calendar
    last_day = calendar.monthrange(year, month)[1]
    end_date = datetime(year, month, last_day).date()

    stmt = select(WorkSchedule).where(WorkSchedule.work_date >= start_date, WorkSchedule.work_date <= end_date)
    result = await db.execute(stmt)
    rows = result.scalars().all()

    schedule_dict = {}
    holiday_list = set()

    for r in rows:
        if r.worker_name not in schedule_dict:
            schedule_dict[r.worker_name] = {}
        schedule_dict[r.worker_name][r.work_date.strftime('%Y-%m-%d')] = r.shift_type

        if r.is_holiday:
            holiday_list.add(r.work_date.strftime('%Y-%m-%d'))

    return {
        "data": schedule_dict,
        "holidays": list(holiday_list)
    }

@router.get("/api/today-workers")
async def get_today_workers(db: AsyncSession = Depends(get_db)):
    today = datetime.now().date()
    today_start = datetime.combine(today, time.min)
    today_end = datetime.combine(today, time.max)
    stmt = select(WorkSchedule).where(
        WorkSchedule.work_date == today,
        or_(
            WorkSchedule.shift_type.like('%○%'),
            WorkSchedule.shift_type.like('%●%'),
            WorkSchedule.shift_type.like('%⚪%'),
            WorkSchedule.shift_type.like('%⚫%')
        )
    )
    result = await db.execute(stmt)
    rows = result.scalars().all()
    leave_stmt = select(Event).where(
        or_(
            Event.color == '#6c757d',
            Event.title.like('%연차%'),
            Event.title.like('%휴가%'),
            Event.title.like('%공가%')
        ),
        or_(
            and_(Event.end == None, Event.start >= today_start, Event.start <= today_end),
            and_(Event.end != None, Event.start <= today_end, Event.end >= today_start)
        )
    )
    leave_result = await db.execute(leave_stmt)
    leave_events = leave_result.scalars().all()
    
    leave_titles = [event.title for event in leave_events if event.title]

    def is_on_leave(name: str) -> bool:
        if not name: return False
        return any(name.strip() in title for title in leave_titles)

    day_workers = [r.worker_name for r in rows if ('○' in r.shift_type or '⚪' in r.shift_type) and not is_on_leave(r.worker_name)]
    night_workers = [r.worker_name for r in rows if ('●' in r.shift_type or '⚫' in r.shift_type) and not is_on_leave(r.worker_name)]
    
    return {
        "date": today.strftime('%Y-%m-%d'),
        "day": day_workers,
        "night": night_workers
    }